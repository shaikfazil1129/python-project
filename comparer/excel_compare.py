import os
import pandas as pd
import logging
import openpyxl
from openpyxl.styles import PatternFill, Font
from .utils import sha256_of_file

def compare_excel_files(file_pairs, target_folder, failed_files_count):
    for fileA, fileB, batch in file_pairs:
        logging.info(f"Comparing batch: {batch} | FileA: {fileA} | FileB: {fileB}")
        sha_A = sha256_of_file(fileA)
        sha_B = sha256_of_file(fileB)
        logging.info(f"SHA-256 for {fileA}: {sha_A}")
        logging.info(f"SHA-256 for {fileB}: {sha_B}")

        if sha_A == sha_B:
            logging.info(f"Skipping {fileA} and {fileB}: files are identical (SHA-256 match).")
            continue

        base_name = f"comparison-{batch}.xlsx"
        outputFile = os.path.join(target_folder, base_name)

        # --- Read both Excel files once ---
        try:
            xlsA = pd.ExcelFile(fileA)
            xlsB = pd.ExcelFile(fileB)
            dataA_all = {sheet: xlsA.parse(sheet_name=sheet) for sheet in xlsA.sheet_names}
            dataB_all = {sheet: xlsB.parse(sheet_name=sheet) for sheet in xlsB.sheet_names}
        except Exception as e:
            logging.error(f"Error reading files {fileA} or {fileB}: {e}")
            failed_files_count += 1
            continue

        sheetsA = xlsA.sheet_names
        sheetsB = xlsB.sheet_names
        max_sheets = min(len(sheetsA), len(sheetsB))
        sheet_diff_map = {}
        output_data = {}

        # --- Compare all sheets first ---
        for index in range(max_sheets):
            sheetA_name = sheetsA[index]
            sheetB_name = sheetsB[index]
            logging.info(f"Comparing sheets: {sheetA_name} vs {sheetB_name}")

            dataA = dataA_all[sheetA_name]
            dataB = dataB_all[sheetB_name]

            if dataA.empty or dataB.empty:
                logging.warning(f"Sheet {sheetA_name} or {sheetB_name} is empty. Skipping.")
                continue

            # Align headers
            missing_headers = [header for header in dataA.columns if header not in dataB.columns]
            for header in missing_headers:
                dataB[header] = pd.NA

            row_count = max(len(dataA), len(dataB))
            output_rows = []
            cell_differences = {}

            for i in range(row_count):
                rowA = dataA.iloc[i] if i < len(dataA) else pd.Series([None] * len(dataA.columns), index=dataA.columns)
                rowB = dataB.iloc[i] if i < len(dataB) else pd.Series([None] * len(dataB.columns), index=dataB.columns)
                diff_details = []
                diff_cols = []
                row_dict = {}

                for header in dataA.columns:
                    valA = rowA[header]
                    valB = rowB[header]
                    row_dict[header] = valB

                    if (pd.isna(valA) and pd.isna(valB)):
                        continue
                    elif pd.isna(valA) or pd.isna(valB) or valA != valB:
                        diff_details.append(f"{header}: A='{valA}', B='{valB}'")
                        diff_cols.append(header)

                if diff_details:
                    cell_differences[i] = diff_cols
                row_dict["Difference"] = "; ".join(diff_details) if diff_details else None
                output_rows.append(row_dict)

            if output_rows and cell_differences:
                df_output = pd.DataFrame(output_rows)
                output_data[sheetB_name] = df_output
                sheet_diff_map[sheetB_name] = {
                    'CellDifferences': cell_differences,
                    'Headers': dataA.columns.tolist()
                }
            else:
                logging.info(f"No mismatches found in sheet {sheetA_name} vs {sheetB_name}, skipping output.")

        # --- Write all output sheets in one go ---
        if output_data:
            file_exists = os.path.exists(outputFile)
            if file_exists:
                try:
                    wb = openpyxl.load_workbook(outputFile)
                except Exception:
                    logging.warning(f"{outputFile} is not a valid Excel file. Recreating.")
                    os.remove(outputFile)
                    file_exists = False
                    wb = None
            else:
                wb = None

            with pd.ExcelWriter(outputFile, engine='openpyxl', mode='a' if file_exists else 'w') as writer:
                for sheet_name, df_output in output_data.items():
                    if file_exists and sheet_name in writer.book.sheetnames:
                        std = writer.book[sheet_name]
                        writer.book.remove(std)
                    df_output.to_excel(writer, sheet_name=sheet_name, index=False)
            logging.info(f"Written all differences to {outputFile}")

        # --- Apply highlights once per file ---
        if sheet_diff_map:
            wb = openpyxl.load_workbook(outputFile)
            for sheet_name, sheet_info in sheet_diff_map.items():
                ws = wb[sheet_name]
                headers = sheet_info['Headers']
                cell_differences = sheet_info['CellDifferences']
                col_map = {header: idx + 1 for idx, header in enumerate(headers)}
                for row_index, cols_changed in cell_differences.items():
                    for col_name in cols_changed:
                        excel_row = row_index + 2
                        excel_col = col_map[col_name]
                        cell = ws.cell(row=excel_row, column=excel_col)
                        cell.fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
                        cell.font = Font(bold=True)
            wb.save(outputFile)
            logging.info(f"Highlighted differences in {outputFile}")
